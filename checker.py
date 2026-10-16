"""Count lead rows by UTM Campaign, then comment and other UTM fields."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
from urllib.parse import unquote
from zipfile import ZipFile


MAX_BYTES = 15 * 1024 * 1024
HEADERS = {"комментарий", "комментарии", "comment", "comments"}
CAMPAIGN_HEADERS = {"utmcampaign", "кампанияutm"}
CAMPAIGN = re.compile(
    r"(?<![\w])utm[_\s-]*campaign\s*[:=]\s*"
    r"(?!utm[_\s-]*(?:source|medium|content|term|id)\s*[:=])"
    r"([^\s&<>\[\]\"';,]+)", re.IGNORECASE,
)
AIKOL_PREFIX = re.compile(r"^(?:айколь?|aikol)(?=$|[\W_])", re.IGNORECASE)
ARCHA_PREFIX = re.compile(r"^(?:арча|archa)(?=$|[\W_])", re.IGNORECASE)


@dataclass
class Result:
    counts: Counter = field(default_factory=Counter)
    total: int = 0
    missing: int = 0
    ambiguous: int = 0
    sheets: list[str] = field(default_factory=list)
    sources: Counter = field(default_factory=Counter)

    def text(self) -> str:
        lines = [f"Всего строк: {self.total}", f"Кампаний: {len(self.counts)}", ""]
        lines.extend(f"{name}: {count}" for name, count in
                     sorted(self.counts.items(), key=lambda item: (-item[1], item[0])))
        lines += ["", f"Без utm_campaign: {self.missing}",
                  f"Несколько разных utm_campaign в строке: {self.ambiguous}", "",
                  f"Из столбца UTM Campaign: {self.sources['column']}",
                  f"Из комментария: {self.sources['comment']}",
                  f"Из других полей: {self.sources['other']}"]
        assert sum(self.counts.values()) + self.missing + self.ambiguous == self.total
        aikol = archa = intersections = 0
        other = []
        for name, count in sorted(self.counts.items(), key=lambda item: (-item[1], item[0])):
            is_aikol = bool(AIKOL_PREFIX.match(name.strip()))
            is_archa = bool(ARCHA_PREFIX.match(name.strip()))
            if is_aikol:
                aikol += count
            if is_archa:
                archa += count
            if is_aikol and is_archa:
                intersections += count
            if not is_aikol and not is_archa:
                other.append((name, count))
        other_total = sum(count for _, count in other) + self.missing + self.ambiguous
        assert aikol + archa - intersections + other_total == self.total
        lines += ["", f"Итого: {aikol + archa - intersections} лидов Айкол / Арча.",
                  f"• Айкол / Aikol: {aikol}", f"• Арча / Archa: {archa}",
                  f"• Пересечений: {intersections}", "", f"Остальные {other_total} лидов:"]
        lines.extend(f"• {name}: {count}" for name, count in other)
        if self.missing:
            lines.append(f"• Без utm_campaign: {self.missing}")
        if self.ambiguous:
            lines.append(f"• Несколько разных кампаний: {self.ambiguous}")
        if not other_total:
            lines.append("• Нет")
        return "\n".join(lines)


def campaign_values(comment: object) -> set[str]:
    return {unquote(match.group(1)) for match in CAMPAIGN.finditer(str(comment or ""))}


def consume(rows, sheet: str, result: Result) -> None:
    column = None
    campaign_columns = []
    header_found = False
    for number, row in enumerate(rows):
        row = list(row)
        labels = [str(value or "").strip().casefold() for value in row]
        candidates = [i for i, value in enumerate(labels) if value in HEADERS]
        field_candidates = [i for i, value in enumerate(labels)
                            if re.sub(r"[\s_-]+", "", value) in CAMPAIGN_HEADERS]
        if not header_found:
            if candidates or field_candidates:
                if len(candidates) > 1:
                    raise ValueError(f"Лист «{sheet}»: несколько столбцов комментариев.")
                column = candidates[0] if candidates else None
                campaign_columns = field_candidates
                header_found = True
                result.sheets.append(sheet)
            elif number >= 49:
                break
            continue
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if (candidates and column in candidates) or (field_candidates and field_candidates == campaign_columns):
            continue  # repeated export header
        result.total += 1
        source = "column"
        values = {str(row[i]).strip() for i in campaign_columns
                  if i < len(row) and row[i] is not None and str(row[i]).strip()}
        if not values:
            source = "comment"
            values = campaign_values(row[column] if column is not None and column < len(row) else None)
        if not values:
            source = "other"
            values = set().union(*(campaign_values(value) for i, value in enumerate(row)
                                   if i != column and i not in campaign_columns))
        if not values:
            result.missing += 1
        elif len(values) > 1:
            result.ambiguous += 1
        else:
            result.counts.update(values)
            result.sources[source] += 1


def analyze(data: bytes) -> Result:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Файл пустой или больше 15 МБ.")
    result = Result()
    if data.startswith(b"PK"):
        from openpyxl import load_workbook
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Распакованный Excel больше 100 МБ.")
        book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book:
                consume(sheet.iter_rows(values_only=True), sheet.title, result)
        finally:
            book.close()
    elif data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        import xlrd
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet in book.sheets():
                consume((sheet.row_values(i) for i in range(sheet.nrows)), sheet.name, result)
        finally:
            book.release_resources()
    elif re.search(br"<table\b", data[:65536], re.IGNORECASE):
        from lxml import html
        parser = html.HTMLParser(no_network=True)
        root = html.fromstring(data, parser=parser)
        tables = root.xpath("//table[not(ancestor::table)]")
        for index, table in enumerate(tables, 1):
            def table_rows():
                for row in table.xpath(".//tr"):
                    if row.xpath("ancestor::table[1]")[0] is table:
                        yield [cell.text_content() for cell in row.xpath("./th|./td")]
            consume(table_rows(), f"Таблица {index}", result)
    else:
        raise ValueError("Нужен .xls, .xlsx или HTML-экспорт Excel из Bitrix.")
    if not result.sheets:
        raise ValueError("Столбец «UTM Campaign» или «Комментарий» не найден в первых 50 строках.")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    args = parser.parse_args()
    try:
        if args.file.stat().st_size > MAX_BYTES:
            raise ValueError("Файл больше 15 МБ.")
        print(analyze(args.file.read_bytes()).text())
    except (ValueError, OSError) as error:
        parser.exit(1, f"Ошибка: {error}\n")


if __name__ == "__main__":
    main()
